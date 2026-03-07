"""Shared content rendering for file-based canary deployers."""

from __future__ import annotations

import os
from io import BytesIO


def render_file_content(text: str, filename: str) -> tuple[bytes, str]:
    """Render canary content as bytes with appropriate content type.

    Returns (data_bytes, content_type) based on filename extension.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".docx":
        return _render_docx(text)
    if ext == ".xlsx":
        return _render_xlsx(text)
    return text.encode("utf-8"), "text/plain; charset=utf-8"


def _render_docx(text: str) -> tuple[bytes, str]:
    from docx import Document

    doc = Document()
    for paragraph in text.split("\n\n"):
        doc.add_paragraph(paragraph.strip())
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _render_xlsx(text: str) -> tuple[bytes, str]:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, line in enumerate(text.strip().splitlines(), start=1):
        ws.cell(row=i, column=1, value=line)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
